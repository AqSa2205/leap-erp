"""Cash Out on the Delivery Board: where it comes from, and who may follow it.

Cash Out is read from Finance's cash-outflow schedule (CashOutflowRow), and the
figure links to Finance -> Cash Outflow only for users that page lets in.
Everyone else sees the number as plain text rather than a link that bounces.
"""
from datetime import date, timedelta
from decimal import Decimal

from django.db.models import prefetch_related_objects
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import Role, User
from finance.models import CashOutflowRow
from pmo import views as pmo_views
from pmo.progress import board_row, cash_out
from projects.models import Project, ProjectStatus, Region

# A figure distinctive enough that finding it in a page means it is ours.
FIGURE = '12,345.67'


def outflow(project, amount, *stages, part='A1'):
    """One finance outflow row; each stage is (date, pct)."""
    fields = {}
    for i, (when, pct) in enumerate(stages, start=1):
        fields[f'date_{i}'] = when
        fields[f'pct_{i}'] = pct
    return CashOutflowRow.objects.create(
        project=project, part=part, amount=Decimal(amount), **fields)


class CashOutEdgeCaseTests(TestCase):

    def setUp(self):
        region = Region.objects.create(name='CO Arabia', code='COLNA', currency='SAR')
        status = ProjectStatus.objects.create(name='CO Ongoing', category='ongoing')
        self.project = Project.objects.create(
            project_name='CO Edge Project', status=status, region=region)
        self.today = timezone.localdate()

    def test_default_as_of_is_today_and_today_counts(self):
        outflow(self.project, '1000',
                (self.today, Decimal('40')),
                (self.today + timedelta(days=1), Decimal('60')))
        self.assertEqual(cash_out(self.project), Decimal('400.00'))

    def test_a_dated_stage_with_no_percentage_adds_nothing(self):
        outflow(self.project, '1000',
                (date(2026, 1, 1), None),
                (date(2026, 1, 2), Decimal('10')))
        self.assertEqual(cash_out(self.project, as_of=date(2026, 6, 1)), Decimal('100.00'))

    def test_a_row_with_no_stages_adds_nothing(self):
        outflow(self.project, '5000')
        self.assertEqual(cash_out(self.project, as_of=date(2026, 6, 1)), Decimal('0.00'))

    def test_a_zero_amount_row_adds_nothing(self):
        outflow(self.project, '0', (date(2026, 1, 1), Decimal('100')))
        self.assertEqual(cash_out(self.project, as_of=date(2026, 6, 1)), Decimal('0.00'))

    def test_all_six_stages_are_read(self):
        outflow(self.project, '1000',
                *[(date(2026, 1, d), Decimal('10')) for d in range(1, 7)])
        self.assertEqual(cash_out(self.project, as_of=date(2026, 6, 1)), Decimal('600.00'))

    def test_prefetched_rows_cost_no_queries(self):
        """board_rows prefetches cash_outflow_rows; cash_out must use that
        cache, or the board runs one query per project."""
        outflow(self.project, '1000', (date(2026, 1, 1), Decimal('100')))
        prefetch_related_objects([self.project], 'cash_outflow_rows')
        with self.assertNumQueries(0):
            cash_out(self.project, as_of=date(2026, 6, 1))

    def test_no_budget_is_never_reported_as_over_budget(self):
        """Cash going out on a project with no approved budget is 'No budget',
        not 'Over Budget' - zero budget and no budget are different things."""
        outflow(self.project, '1000', (date(2020, 1, 1), Decimal('100')))
        row = board_row(self.project)
        self.assertEqual(row['cash_out'], Decimal('1000.00'))
        self.assertEqual(row['budget_status'], 'none')


class BoardCashOutLinkAccessTests(TestCase):

    def setUp(self):
        self.roles = {
            name: Role.objects.get_or_create(name=name)[0]
            for name in (Role.PROJECT_MANAGER, Role.SUPER_ADMIN, Role.DEVELOPER,
                         Role.FINANCE_HEAD, Role.FINANCE_MANAGER, Role.FINANCE_REP)}
        self.region = Region.objects.create(name='CO Board Arabia', code='COBLNA', currency='SAR')
        self.status = ProjectStatus.objects.create(name='CO Board Ongoing', category='ongoing')
        self.project = Project.objects.create(
            project_name='CO Board Project', status=self.status, region=self.region)
        outflow(self.project, '12345.67', (date(2020, 1, 1), Decimal('100')))

        self.board = reverse('pmo:board')
        self.link = reverse('finance:cash_outflow', args=[self.project.pk])

        self.pm = self.user('co_pm', Role.PROJECT_MANAGER)
        self.admin = self.user('co_admin', Role.SUPER_ADMIN)
        self.dev = self.user('co_dev', Role.DEVELOPER)
        self.finance_users = [
            self.user('co_fin_head', Role.FINANCE_HEAD),
            self.user('co_fin_mgr', Role.FINANCE_MANAGER),
            self.user('co_fin_rep', Role.FINANCE_REP)]

    def user(self, username, role_name, region=None):
        return User.objects.create_user(
            username, password='pw', role=self.roles[role_name],
            region=region or self.region)

    def get_board(self, user):
        self.client.force_login(user)
        return self.client.get(self.board)

    # --- who gets the link -------------------------------------------------

    def test_pm_sees_the_figure_as_plain_text(self):
        r = self.get_board(self.pm)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'CO Board Project')
        self.assertContains(r, FIGURE)
        self.assertNotContains(r, self.link)
        self.assertFalse(r.context['can_open_finance'])

    def test_the_link_hidden_from_a_pm_really_would_bounce_them(self):
        """Pins why the link is hidden: Finance -> Cash Outflow turns a PM
        away. If finance ever opens that page to PMs, this fails and the
        board can show them the link too."""
        self.client.force_login(self.pm)
        r = self.client.get(self.link)
        self.assertRedirects(r, reverse('dashboard:index'), fetch_redirect_response=False)

    def test_super_admin_gets_the_link(self):
        r = self.get_board(self.admin)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, self.link)
        self.assertContains(r, FIGURE)
        self.assertTrue(r.context['can_open_finance'])

    def test_finance_roles_get_the_link_or_are_kept_off_the_board(self):
        """Every finance role either sees a working link or cannot open the
        board at all - never a page showing cash figures without it."""
        for u in self.finance_users:
            with self.subTest(role=u.role.name):
                self.client.force_login(u)
                r = self.client.get(self.board)
                if pmo_views.can_see_delivery(u):
                    self.assertEqual(r.status_code, 200)
                    self.assertContains(r, self.link)
                else:
                    self.assertEqual(r.status_code, 403)
                    self.assertNotContains(r, FIGURE, status_code=403)

    # --- who sees the figures at all ----------------------------------------

    def test_role_without_delivery_access_sees_no_cash_figures(self):
        r = self.get_board(self.dev)
        self.assertEqual(r.status_code, 403)
        self.assertNotContains(r, FIGURE, status_code=403)

    def test_anonymous_is_sent_to_login(self):
        r = self.client.get(self.board)
        self.assertEqual(r.status_code, 302)
        self.assertIn('login', r.url)

    def test_pm_in_another_region_does_not_see_this_projects_cash_out(self):
        other = Region.objects.create(name='CO Other Region', code='COOTHER', currency='SAR')
        outsider = self.user('co_outsider', Role.PROJECT_MANAGER, region=other)
        r = self.get_board(outsider)
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, 'CO Board Project')
        self.assertNotContains(r, FIGURE)


class BoardExportCashOutTests(TestCase):
    """The Excel export carries the same Cash Out as the board, and its
    note no longer says Cash Out is committed purchase orders."""

    def setUp(self):
        pm_role, _ = Role.objects.get_or_create(name=Role.PROJECT_MANAGER)
        dev_role, _ = Role.objects.get_or_create(name=Role.DEVELOPER)
        region = Region.objects.create(name='CO Export Arabia', code='COXLNA', currency='SAR')
        status = ProjectStatus.objects.create(name='CO Export Ongoing', category='ongoing')
        self.project = Project.objects.create(
            project_name='CO Export Project', status=status, region=region)
        outflow(self.project, '12345.67', (date(2020, 1, 1), Decimal('100')))
        outflow(self.project, '50000', (date(2999, 1, 1), Decimal('100')))   # future
        self.pm = User.objects.create_user('co_x_pm', password='pw', role=pm_role, region=region)
        self.dev = User.objects.create_user('co_x_dev', password='pw', role=dev_role, region=region)
        self.url = reverse('pmo:board_export_excel')

    def workbook(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.client.force_login(self.pm)
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        return load_workbook(BytesIO(r.content))['Budget']

    def test_cash_out_column_is_finance_paid_to_date_ex_vat(self):
        ws = self.workbook()
        rows = list(ws.iter_rows(values_only=True))
        header = next(row for row in rows
                      if any(str(v or '').startswith('Cash Out') for v in row))
        col = next(i for i, v in enumerate(header) if str(v or '').startswith('Cash Out'))
        project_row = next(row for row in rows
                           if any('CO Export Project' in str(v or '') for v in row))
        self.assertEqual(Decimal(str(project_row[col])).quantize(Decimal('0.01')),
                         Decimal('12345.67'))

    def test_note_describes_the_finance_source(self):
        note = str(self.workbook()['A2'].value)
        self.assertIn('cash outflow schedule', note)
        self.assertNotIn('committed purchase orders', note)

    def test_role_without_delivery_access_cannot_export(self):
        self.client.force_login(self.dev)
        self.assertEqual(self.client.get(self.url).status_code, 403)
