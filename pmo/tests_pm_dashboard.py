import json
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from accounts.models import Role, User
from costing.models import CostingSheet
from projects.models import Document, Project, ProjectStatus, Region

from .models import ResponsibilityMatrix


class PmDashboardAccessTests(TestCase):
    """The PM Dashboard is gated by pmo.views.can_see_delivery — the same
    role ladder as the Delivery Board, since it's the same Project
    Management audience looking at different information about the same
    projects."""

    def setUp(self):
        self.pm_role, _ = Role.objects.get_or_create(name=Role.PROJECT_MANAGER)
        self.dev_role, _ = Role.objects.get_or_create(name=Role.DEVELOPER)
        self.region = Region.objects.create(name='Arabia', code='LNA', currency='SAR')
        self.status = ProjectStatus.objects.create(name='Won', category='won')
        self.project = Project.objects.create(
            project_name='Test Project', status=self.status, region=self.region,
            po_number='PO-123')
        self.pm_user = User.objects.create_user(
            'pm1', password='pw', role=self.pm_role, region=self.region)
        self.dev_user = User.objects.create_user(
            'dev1', password='pw', role=self.dev_role, region=self.region)

    def test_role_without_delivery_access_is_denied(self):
        self.client.force_login(self.dev_user)
        r = self.client.get(reverse('pmo:pm_dashboard_index'))
        self.assertEqual(r.status_code, 403)

    def test_project_manager_can_open_portfolio_and_project_overview(self):
        self.client.force_login(self.pm_user)
        self.assertEqual(self.client.get(reverse('pmo:pm_dashboard_index')).status_code, 200)
        r = self.client.get(reverse('pmo:pm_project_overview', args=[self.project.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, self.project.project_name)

    def test_project_po_page_shows_po_number(self):
        self.client.force_login(self.pm_user)
        r = self.client.get(reverse('pmo:pm_project_po', args=[self.project.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'PO-123')

    def test_export_excel_returns_a_workbook(self):
        import openpyxl
        from io import BytesIO

        self.client.force_login(self.pm_user)
        r = self.client.get(reverse('pmo:pm_dashboard_export_excel'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        wb = openpyxl.load_workbook(BytesIO(r.content))
        ws = wb.active
        header_row = [c.value for c in ws[4]]
        self.assertEqual(header_row[0], 'Project Name')
        self.assertIn('Test Project', [row[0].value for row in ws.iter_rows(min_row=5)])


class ResponsibilityMatrixEditTests(TestCase):
    def setUp(self):
        self.pm_role, _ = Role.objects.get_or_create(name=Role.PROJECT_MANAGER)
        self.region = Region.objects.create(name='Arabia', code='LNA', currency='SAR')
        self.status = ProjectStatus.objects.create(name='Won', category='won')
        self.project = Project.objects.create(
            project_name='Test Project', status=self.status, region=self.region)
        self.pm_user = User.objects.create_user(
            'pm2', password='pw', role=self.pm_role, region=self.region)

    def test_edit_saves_sanitized_grid_and_drops_empty_rows(self):
        self.client.force_login(self.pm_user)
        url = reverse('pmo:pm_responsibility_matrix_edit', args=[self.project.pk])
        payload = {
            'columns_json': json.dumps([{'key': 'c1', 'name': 'Task'}, {'key': 'c2', 'name': 'Owner'}]),
            'rows_json': json.dumps([
                {'cells': {'c1': {'text': 'Kickoff'}, 'c2': {'text': 'PM'}}},
                {'cells': {'c1': {'text': ''}, 'c2': {'text': ''}}},  # all-empty row, should be dropped
            ]),
        }
        r = self.client.post(url, payload)
        self.assertRedirects(r, reverse('pmo:pm_responsibility_matrix', args=[self.project.pk]))
        matrix = ResponsibilityMatrix.objects.get(project=self.project)
        self.assertEqual(len(matrix.rows), 1)
        self.assertEqual(matrix.rows[0]['cells']['c1']['text'], 'Kickoff')
        self.assertEqual(matrix.updated_by, self.pm_user)


class CommunicationMatrixExportTests(TestCase):
    def setUp(self):
        self.pm_role, _ = Role.objects.get_or_create(name=Role.PROJECT_MANAGER)
        self.region = Region.objects.create(name='Arabia', code='LNA', currency='SAR')
        self.status = ProjectStatus.objects.create(name='Won', category='won')
        self.project = Project.objects.create(
            project_name='Test Project', status=self.status, region=self.region)
        self.pm_user = User.objects.create_user(
            'pm3', password='pw', role=self.pm_role, region=self.region)

    def test_export_pdf(self):
        self.client.force_login(self.pm_user)
        url = reverse('pmo:pm_communication_matrix_export_pdf', args=[self.project.pk])
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')


class DataIsolationTests(TestCase):
    """The PO documents and BOQ (costing sheets) shown for a project come
    from that project's own FK relations only (Document.project,
    CostingSheet.project) — never from a title match or a global pool."""

    def setUp(self):
        self.pm_role, _ = Role.objects.get_or_create(name=Role.PROJECT_MANAGER)
        self.region = Region.objects.create(name='Arabia', code='LNA', currency='SAR')
        self.status = ProjectStatus.objects.create(name='Won', category='won')
        self.pm_user = User.objects.create_user(
            'pm4', password='pw', role=self.pm_role, region=self.region)

        self.project_a = Project.objects.create(
            project_name='Project Alpha', status=self.status, region=self.region,
            proposal_reference='REF-ALPHA', po_number='PO-ALPHA-001')
        self.project_b = Project.objects.create(
            project_name='Project Beta', status=self.status, region=self.region,
            proposal_reference='REF-BETA', po_number='PO-BETA-002')

        Document.objects.create(
            project=self.project_a, name='Alpha PO Document', document_type='po_document',
            uploaded_by=self.pm_user,
            file=SimpleUploadedFile('alpha_po.pdf', b'dummy-alpha', content_type='application/pdf'))
        Document.objects.create(
            project=self.project_b, name='Beta PO Document', document_type='po_document',
            uploaded_by=self.pm_user,
            file=SimpleUploadedFile('beta_po.pdf', b'dummy-beta', content_type='application/pdf'))

        CostingSheet.objects.create(
            project=self.project_a, title='Alpha BOQ Sheet',
            margin=Decimal('25'), discount_rate=Decimal('0'))
        CostingSheet.objects.create(
            project=self.project_b, title='Beta BOQ Sheet',
            margin=Decimal('25'), discount_rate=Decimal('0'))

    def test_po_page_only_shows_this_projects_document(self):
        self.client.force_login(self.pm_user)
        r_a = self.client.get(reverse('pmo:pm_project_po', args=[self.project_a.pk]))
        self.assertContains(r_a, 'Alpha PO Document')
        self.assertNotContains(r_a, 'Beta PO Document')

        r_b = self.client.get(reverse('pmo:pm_project_po', args=[self.project_b.pk]))
        self.assertContains(r_b, 'Beta PO Document')
        self.assertNotContains(r_b, 'Alpha PO Document')

    def test_boq_page_only_shows_this_projects_costing_sheet(self):
        self.client.force_login(self.pm_user)
        r_a = self.client.get(reverse('pmo:pm_project_boq', args=[self.project_a.pk]))
        self.assertContains(r_a, 'Alpha BOQ Sheet')
        self.assertNotContains(r_a, 'Beta BOQ Sheet')

        r_b = self.client.get(reverse('pmo:pm_project_boq', args=[self.project_b.pk]))
        self.assertContains(r_b, 'Beta BOQ Sheet')
        self.assertNotContains(r_b, 'Alpha BOQ Sheet')

    def test_overview_counts_do_not_leak_across_projects(self):
        self.client.force_login(self.pm_user)
        r_a = self.client.get(reverse('pmo:pm_project_overview', args=[self.project_a.pk]))
        self.assertEqual(r_a.status_code, 200)
        self.assertContains(r_a, 'Project Alpha')
        self.assertNotContains(r_a, 'Project Beta')
