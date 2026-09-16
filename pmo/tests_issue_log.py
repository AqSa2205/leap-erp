from django.test import TestCase
from django.urls import reverse

from accounts.models import Role, User
from projects.models import Project, ProjectStatus, Region

from .models import ProjectIssue


class IssueLogAccessTests(TestCase):
    def setUp(self):
        self.pm_role, _ = Role.objects.get_or_create(name=Role.PROJECT_MANAGER)
        self.manager_role, _ = Role.objects.get_or_create(name=Role.MANAGER)
        self.dev_role, _ = Role.objects.get_or_create(name=Role.DEVELOPER)
        self.pm_user = User.objects.create_user('il_pm', password='pw', role=self.pm_role)
        self.manager_user = User.objects.create_user('il_mgr', password='pw', role=self.manager_role)
        self.dev_user = User.objects.create_user('il_dev', password='pw', role=self.dev_role)

        self.region = Region.objects.create(name='Arabia', code='LNA', currency='SAR')
        self.status = ProjectStatus.objects.create(name='Ongoing', category='ongoing')
        self.project = Project.objects.create(
            project_name='Test Project', status=self.status, region=self.region)
        # projects_visible_to falls back to "owned only" for a role with no
        # region — give the PM a matching region so the project shows up in
        # their scoped choices, same as a real project_manager user would.
        self.pm_user.region = self.region
        self.pm_user.save()

    def test_role_without_delivery_access_is_denied(self):
        self.client.force_login(self.dev_user)
        self.assertEqual(self.client.get(reverse('pmo:issue_log_list')).status_code, 403)

    def test_manager_can_view_but_not_log(self):
        self.client.force_login(self.manager_user)
        self.assertEqual(self.client.get(reverse('pmo:issue_log_list')).status_code, 200)
        self.assertEqual(self.client.get(reverse('pmo:issue_log_create')).status_code, 403)

    def test_project_manager_can_log_an_issue(self):
        self.client.force_login(self.pm_user)
        r = self.client.post(reverse('pmo:issue_log_create'), {
            'project': self.project.pk,
            'status': 'open',
            'priority': 'critical',
            'description': 'Something broke.',
            'owner': 'LEAP',
            'date_identified': '2026-08-05',
            'estimated_resolution_date': '2026-08-20',
            'escalation_needed': 'on',
            'impact': 'Big impact.',
            'actions': 'Fix it.',
            'actual_resolution_date': '',
            'final_resolution': '',
        })
        self.assertRedirects(r, reverse('pmo:issue_log_list') + '?year=2026&month=8')
        issue = ProjectIssue.objects.get(project=self.project)
        self.assertEqual(issue.logged_by, self.pm_user)
        self.assertTrue(issue.escalation_needed)
        self.assertEqual(issue.priority_color, ProjectIssue.PRIORITY_COLORS['critical'])

    def test_list_only_shows_issues_for_the_selected_month(self):
        ProjectIssue.objects.create(
            project=self.project, description='August issue',
            date_identified='2026-08-05', logged_by=self.pm_user)
        ProjectIssue.objects.create(
            project=self.project, description='September issue',
            date_identified='2026-09-05', logged_by=self.pm_user)
        self.client.force_login(self.pm_user)

        r_aug = self.client.get(reverse('pmo:issue_log_list') + '?year=2026&month=8')
        self.assertContains(r_aug, 'August issue')
        self.assertNotContains(r_aug, 'September issue')

        r_sep = self.client.get(reverse('pmo:issue_log_list') + '?year=2026&month=9')
        self.assertContains(r_sep, 'September issue')
        self.assertNotContains(r_sep, 'August issue')

    def test_export_excel_matches_the_source_workbook_columns(self):
        import openpyxl
        from io import BytesIO

        ProjectIssue.objects.create(
            project=self.project, description='Exported issue', owner='Aramco',
            date_identified='2026-08-05', logged_by=self.pm_user)
        self.client.force_login(self.pm_user)
        r = self.client.get(reverse('pmo:issue_log_export_excel') + '?year=2026&month=8')
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(r.content))
        ws = wb.active
        # Row 1: title. Rows 2-5: the 4-row severity legend. Row 6: spacer.
        # Row 7: header. Row 8+: data.
        header_row = [c.value for c in ws[7]]
        self.assertEqual(header_row[:5], [
            'ID', 'Status', 'Priority', 'Description', 'Project Name'])
        data_row = [c.value for c in ws[8]]
        self.assertEqual(data_row[3], 'Exported issue')
        self.assertEqual(data_row[5], 'Aramco')

    def test_export_all_zip_contains_one_workbook_per_month(self):
        ProjectIssue.objects.create(
            project=self.project, description='August issue',
            date_identified='2026-08-05', logged_by=self.pm_user)
        ProjectIssue.objects.create(
            project=self.project, description='September issue',
            date_identified='2026-09-05', logged_by=self.pm_user)
        self.client.force_login(self.pm_user)
        r = self.client.get(reverse('pmo:issue_log_export_all_zip'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/zip')

        import zipfile
        from io import BytesIO
        with zipfile.ZipFile(BytesIO(r.content)) as zf:
            names = zf.namelist()
        self.assertEqual(len(names), 2)
        self.assertTrue(any('August' in n for n in names))
        self.assertTrue(any('September' in n for n in names))
