"""Project delivery — the overview board and one project's milestone WBS."""
import calendar
import json
import zipfile
from decimal import Decimal, InvalidOperation
from functools import wraps
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core import signing
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Count, Prefetch, Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from dashboard.views import projects_visible_to
from projects.models import Project

from .forms import ManpowerDetailsForm, NewEmployeeManpowerForm, ProjectIssueForm
from .models import (ONE, ZERO, CommunicationMatrix, ManpowerResource,
                      MilestoneProgressEntry, ProjectIssue, ProjectMilestone,
                      ResponsibilityMatrix, default_communication_columns,
                      default_responsibility_columns, sanitize_grid)
from .progress import (board_row, board_rows, leaves, milestone_checklist,
                       project_completion, validate_weightages)
from .workbook_import import (deserialise_activities, plan_workbook,
                              serialise_activities, write_activities)


def can_see_delivery(user):
    """Who gets the Project Management department.

    Project Manager and Site Manager are the delivery roles; the rest is the
    usual oversight ladder. Kept as one function so the sidebar, the board and
    the update endpoint cannot disagree about who is allowed in — the workbook
    equivalent was a shared drive with no answer to this at all.
    """
    if not getattr(user, 'is_authenticated', False):
        return False
    return bool(
        user.is_super_admin_user
        or user.is_admin_user
        or user.is_manager_user
        or user.is_project_manager_user
        or user.is_site_manager_user
    )


def can_update_progress(user):
    """Who may move a completion figure.

    Narrower than viewing on purpose: the whole company's delivery reporting
    is derived from these numbers, so read access and write access are not the
    same decision.
    """
    if not getattr(user, 'is_authenticated', False):
        return False
    return bool(
        user.is_super_admin_user
        or user.is_admin_user
        or user.is_project_manager_user
        or user.is_site_manager_user
    )


def delivery_required(view_func):
    """Gate a view behind can_see_delivery.

    The PM Dashboard, Manpower Status and Issue Log views all share this one
    check with the same message — pulled into a decorator instead of
    hand-repeating `if not can_see_delivery(...): raise PermissionDenied(...)`
    at the top of every one of them, which is easy to forget or reword
    slightly differently at any single call site.
    """
    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not can_see_delivery(request.user):
            raise PermissionDenied('The Project Management department is not open to your role.')
        return view_func(request, *args, **kwargs)
    return _wrapped


def update_access_required(message):
    """Gate a view behind the narrower can_update_progress, with a message
    specific to what that view lets someone change."""
    def _decorator(view_func):
        @wraps(view_func)
        def _wrapped(request, *args, **kwargs):
            if not can_update_progress(request.user):
                raise PermissionDenied(message)
            return view_func(request, *args, **kwargs)
        return _wrapped
    return _decorator


def _visible_projects(user):
    """Delivery projects this user may see, with the milestone tree attached.

    Scoping reuses `dashboard.views.projects_visible_to` rather than growing a
    third copy of the same ladder — that function exists because two copies
    had already drifted apart.
    """
    return (
        projects_visible_to(user)
        .select_related('region', 'status', 'finance')
        .prefetch_related(
            Prefetch('milestones',
                     queryset=ProjectMilestone.objects.prefetch_related('progress_entries')),
            'purchase_orders',
            'finance__milestones',
        )
    )


@login_required
def board(request):
    """Every project on one page, every column derived.

    Replaces the Projects Overview sheet, whose totals summed rows 7–13 while
    the data ran to row 17. There is no range here to get wrong.
    """
    if not can_see_delivery(request.user):
        raise PermissionDenied('The Project Management department is not open to your role.')

    projects = list(_visible_projects(request.user))
    rows = board_rows(projects)
    total_in = sum((r['cash_in'] for r in rows), ZERO)
    total_out = sum((r['cash_out'] for r in rows), ZERO)

    return render(request, 'pmo/board.html', {
        'rows': rows,
        'total_cash_in': total_in,
        'total_cash_out': total_out,
        'total_net': total_in - total_out,
        'problem_count': sum(1 for r in rows if r['weightage_problems']),
        'can_update': can_update_progress(request.user),
    })


@login_required
def project_detail(request, pk):
    """One project's WBS, parents with their children underneath."""
    if not can_see_delivery(request.user):
        raise PermissionDenied('The Project Management department is not open to your role.')

    project = get_object_or_404(_visible_projects(request.user), pk=pk)
    rows = list(project.milestones.all())
    children = {}
    for row in rows:
        if row.parent_id is not None:
            children.setdefault(row.parent_id, []).append(row)

    # Flattened for the template: a parent immediately followed by its own
    # children, so the grid is one <tbody> and keyboard navigation runs down it
    # in the order somebody reads.
    #
    # Weights are summed here rather than read from the model properties. Those
    # walk `self.children` per row, which is a query each — fine for one row on
    # a page, but this page is the whole tree.
    display = []
    for parent in sorted((r for r in rows if r.parent_id is None), key=lambda r: r.order):
        kids = sorted(children.get(parent.pk, []), key=lambda r: r.order)
        display.append({
            'row': parent,
            'is_parent': True,
            'number': str(parent.order),
            'weightage': sum((k.weightage for k in kids), ZERO) if kids else parent.weightage,
            'completed_weightage': sum(
                (k.weightage * k.completed_fraction for k in kids), ZERO),
        })
        for child in kids:
            display.append({
                'row': child,
                'is_parent': False,
                'number': f'{parent.order}.{child.order}',
                'weightage': child.weightage,
                'completed_weightage': child.weightage * child.completed_fraction,
            })

    completion = project_completion(project)
    return render(request, 'pmo/project_detail.html', {
        'project': project,
        'display': display,
        'completion': completion,
        'completion_pct': completion * Decimal('100'),
        'problems': validate_weightages(project),
        'leaf_count': len(leaves(project)),
        'can_update': can_update_progress(request.user),
    })


@require_POST
@login_required
def update_progress(request, pk):
    """Set one milestone's completed fraction. Called per cell by the grid.

    Every change is appended to the progress log as well as written to the
    row, which is what makes "when was this last updated" answerable — the
    workbook's version of that column was TODAY() and always read as today.
    """
    if not can_update_progress(request.user):
        return JsonResponse({'error': 'You cannot update delivery progress.'}, status=403)

    milestone = get_object_or_404(
        ProjectMilestone.objects.select_related('project'), pk=pk)
    if not _visible_projects(request.user).filter(pk=milestone.project_id).exists():
        return JsonResponse({'error': 'Project not found.'}, status=404)

    # Progress belongs to the activities that carry weight. Accepting it on a
    # summary row would let a parent and its children both claim the same
    # weight, and the project would read as more complete than it is.
    if milestone.children.exists():
        return JsonResponse(
            {'error': 'This is a summary row — update the activities under it.'},
            status=400)

    raw = (request.POST.get('completed_fraction') or '').strip()
    try:
        value = Decimal(raw)
    except (InvalidOperation, ValueError):
        return JsonResponse({'error': f'"{raw}" is not a number.'}, status=400)
    if value < ZERO or value > ONE:
        return JsonResponse(
            {'error': 'Progress is a fraction between 0 and 1.'}, status=400)

    milestone.completed_fraction = value
    milestone.save(update_fields=['completed_fraction', 'updated_at'])
    MilestoneProgressEntry.objects.create(
        milestone=milestone, completed_fraction=value,
        note=(request.POST.get('note') or '').strip(),
        recorded_by=request.user)

    project = milestone.project
    return JsonResponse({
        'ok': True,
        'completed_fraction': str(value),
        'completed_weightage': str(milestone.weightage * value),
        'project_completion_pct': str(project_completion(project) * Decimal('100')),
    })


# ── importing the workbook from the browser ─────────────────────────────────
#
# The management command that does this cannot be run in production: the web
# service runs on a plan with no shell at all. So the one-off that actually
# matters — loading ten milestone sheets — has to be doable from a screen, the
# same way finance publishes a chart revision at /accounting/chart/import/.

MAX_WORKBOOK_BYTES = 15 * 1024 * 1024
PREVIEW_SALT = 'pmo.milestone_import'
PREVIEW_MAX_AGE = 60 * 60


def can_import_milestones(user):
    """Importing rewrites whole projects' WBS, so it sits above the weekly
    update rather than beside it: a project manager moves a figure, an
    administrator replaces the structure."""
    if not getattr(user, 'is_authenticated', False):
        return False
    return bool(user.is_super_admin_user or user.is_admin_user)


@login_required
def milestone_import(request):
    """Upload the projects overview workbook, see what it would do, then apply.

    The preview is the whole safety story. Sheets are matched on the Leap
    reference and an unmatched one is reported rather than guessed, so what a
    person has to check before confirming is exactly what this page shows:
    which sheet became which project, which projects already have milestones,
    and which weights do not add up.

    The parsed activities travel to the confirm step in a signed hidden field.
    Re-reading the upload on confirm would mean applying something other than
    what was previewed, which is the bug the chart importer had.
    """
    if not can_import_milestones(request.user):
        raise PermissionDenied('Importing milestones is limited to administrators.')

    if request.method != 'POST':
        return render(request, 'pmo/milestone_import.html', {})

    upload = request.FILES.get('workbook')
    if not upload:
        messages.error(request, 'Choose a workbook to upload.')
        return redirect('pmo:milestone_import')
    if upload.size > MAX_WORKBOOK_BYTES:
        messages.error(request, 'That file is larger than 15 MB, which is almost '
                                'certainly not the projects overview workbook.')
        return redirect('pmo:milestone_import')

    try:
        import openpyxl
        workbook = openpyxl.load_workbook(
            BytesIO(upload.read()), data_only=True, read_only=True)
    except Exception:
        # Deliberately broad: openpyxl raises several unrelated exception types
        # for a file that simply is not a workbook, and every one of them means
        # the same thing to somebody who picked the wrong file.
        messages.error(request, 'That file could not be read as an .xlsx workbook.')
        return redirect('pmo:milestone_import')

    results = plan_workbook(workbook, list(Project.objects.all()))
    if not results:
        messages.error(request, 'No milestone sheets found in that workbook. The '
                                'importer looks for sheets named after a Leap '
                                'reference, such as "LNA-2308-Milestones(MASCO)".')
        return redirect('pmo:milestone_import')

    for result in results:
        project = result['project']
        result['existing'] = project.milestones.count() if project else 0

    # Only matched sheets can be applied, so only those are signed. A sheet
    # with no project is a question for a person, not a thing to carry forward.
    payload = [{
        'project_id': r['project'].pk,
        'sheet': r['sheet'],
        'activities': serialise_activities(r['activities']),
    } for r in results if r['project'] and r['activities']]

    return render(request, 'pmo/milestone_import.html', {
        'results': results,
        'filename': upload.name,
        'matched': [r for r in results if r['project']],
        'unmatched': [r for r in results if not r['project']],
        'with_problems': [r for r in results if r['problems']],
        'already_have': [r for r in results if r['existing']],
        'payload': signing.dumps(payload, salt=PREVIEW_SALT, compress=True),
    })


@require_POST
@login_required
def milestone_import_apply(request):
    """Write the milestones that were previewed."""
    if not can_import_milestones(request.user):
        raise PermissionDenied('Importing milestones is limited to administrators.')

    raw = request.POST.get('payload') or ''
    if not raw:
        messages.error(request, 'Nothing to apply — upload the workbook first.')
        return redirect('pmo:milestone_import')

    try:
        payload = signing.loads(raw, salt=PREVIEW_SALT, max_age=PREVIEW_MAX_AGE)
    except signing.SignatureExpired:
        messages.error(request, 'That preview is more than an hour old and the '
                                'projects may have moved on since. Please upload '
                                'the workbook again.')
        return redirect('pmo:milestone_import')
    except signing.BadSignature:
        messages.error(request, 'That preview could not be verified. Please upload '
                                'the workbook again.')
        return redirect('pmo:milestone_import')

    replace = request.POST.get('replace') == 'on'
    written = skipped = 0
    # One transaction: a half-written import would leave projects with a
    # partial WBS whose weights cannot add up, which reads as a data problem
    # rather than as an import that failed.
    with transaction.atomic():
        for entry in payload:
            project = Project.objects.filter(pk=entry['project_id']).first()
            if project is None:
                continue                      # deleted since the preview
            if project.milestones.exists():
                if not replace:
                    skipped += 1
                    continue
                project.milestones.all().delete()
            written += write_activities(
                project, deserialise_activities(entry['activities']))

    messages.success(request, f'Imported {written} milestone rows.')
    if skipped:
        messages.warning(
            request,
            f'{skipped} project(s) already had milestones and were left alone. '
            f'Tick "replace existing milestones" to overwrite them.')
    return redirect('pmo:board')


# ── PM Dashboard: client PO, unpriced BOQ, responsibility & communication ──
#
# A second, client-facing view onto the same projects `board` covers —
# gated the same way (can_see_delivery) since it's the same Project
# Management audience, just looking at different information about the
# same projects.

def _pm_visible_project_or_404(request, pk):
    return get_object_or_404(_visible_projects(request.user), pk=pk)


@login_required
@delivery_required
def pm_dashboard_index(request):
    """Portfolio landing page: graphical overview of every delivery project
    this user can see, then the searchable table to pick one and drill in."""
    projects_qs = _visible_projects(request.user)

    q = (request.GET.get('q') or '').strip()
    filtered = projects_qs
    if q:
        filtered = filtered.filter(
            Q(project_name__icontains=q)
            | Q(serial_number__icontains=q)
            | Q(customer__icontains=q)
            | Q(po_number__icontains=q)
        )
    filtered = filtered.order_by('-id')
    rows = board_rows(list(filtered))

    total_projects = projects_qs.count()
    total_value = projects_qs.aggregate(v=Sum('estimated_value'))['v'] or 0
    po_missing = projects_qs.filter(Q(po_number='') | Q(po_number__isnull=True)).count()

    status_counts = list(
        projects_qs.exclude(status__isnull=True)
        .values('status__name', 'status__color')
        .annotate(count=Count('id')).order_by('-count')
    )
    region_counts = list(
        projects_qs.exclude(region__isnull=True)
        .values('region__code')
        .annotate(count=Count('id')).order_by('-count')
    )

    milestone_buckets = {'success': 0, 'warning': 0, 'danger': 0, '': 0}
    for p in projects_qs:
        key = p.milestone_overall_status or ''
        milestone_buckets[key] = milestone_buckets.get(key, 0) + 1

    return render(request, 'pmo/pm_dashboard_index.html', {
        'rows': rows,
        'q': q,
        'total_projects': total_projects,
        'total_value': total_value,
        'po_missing': po_missing,
        'milestone_on_track': milestone_buckets.get('success', 0) + milestone_buckets.get('warning', 0),
        'milestone_at_risk': milestone_buckets.get('danger', 0),
        'milestone_no_data': milestone_buckets.get('', 0),
        'status_labels_json': json.dumps([s['status__name'] or 'Unknown' for s in status_counts]),
        'status_data_json': json.dumps([s['count'] for s in status_counts]),
        'status_colors_json': json.dumps([s['status__color'] or '#6c757d' for s in status_counts]),
        'region_labels_json': json.dumps([r['region__code'] or 'Unknown' for r in region_counts]),
        'region_data_json': json.dumps([r['count'] for r in region_counts]),
    })


def _style_export_header_and_grid(ws, header_idx, ncols):
    """Shared openpyxl look for this module's Excel exports: a dark header
    row, a thin grey border on every cell from the header down, and the
    header row frozen. Pulled into one place so pm_dashboard_export_excel
    and _build_issue_log_workbook don't each hand-declare the same
    PatternFill/Font/Border boilerplate."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    header_fill = PatternFill('solid', fgColor='2C3E50')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='D5D5D5')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in ws[header_idx]:
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=header_idx, max_row=ws.max_row, min_col=1, max_col=ncols):
        for c in row:
            c.border = border
    ws.freeze_panes = ws.cell(row=header_idx + 1, column=1)


@login_required
@delivery_required
def pm_dashboard_export_excel(request):
    """The PM Dashboard's project list as an .xlsx, one row per project, in
    the same column order as the PM's own tracking workbook. Honours the
    same search box as the page it's exported from."""
    import openpyxl
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    projects_qs = _visible_projects(request.user)
    q = (request.GET.get('q') or '').strip()
    if q:
        projects_qs = projects_qs.filter(
            Q(project_name__icontains=q)
            | Q(serial_number__icontains=q)
            | Q(customer__icontains=q)
            | Q(po_number__icontains=q)
        )
    rows = board_rows(list(projects_qs.order_by('-id')))

    columns = [
        'Project Name', 'Expected Completion Date', 'Weightage of Pending Activities',
        'Completed Activity Weightage', 'Planned Project End Date', 'Delay (Months)',
        'Cash In (SAR)', 'Achieved Milestones', 'No. Pending Milestones', 'Action By',
        'Project Status',
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'PM Dashboard'
    ncols = len(columns)

    delay_fill = PatternFill('solid', fgColor='FFC7CE')
    green_font = Font(bold=True, color='198754')
    red_font = Font(bold=True, color='DC3545')

    ws.append(['PM Dashboard — Project Status'])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws['A1'].font = Font(bold=True, size=14, color='404040')
    generated = timezone.localtime(timezone.now()).strftime('%d %b %Y at %H:%M')
    ws.append([f'Generated {generated}'])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws['A2'].font = Font(italic=True, size=9, color='666666')
    ws.append([])

    header_idx = ws.max_row + 1
    ws.append(columns)

    for row in rows:
        project = row['project']
        pending_pct = float(100 - row['completion_pct'])
        completed_pct = float(row['completion_pct'])
        ws.append([
            project.project_name,
            row['expected_completion_date'],
            pending_pct / 100,
            completed_pct / 100,
            row['end_date'],
            row['delay_months'] or 0,
            float(row['cash_in']),
            row['achieved_milestones'],
            row['pending_milestones'],
            row['action_by'],
            project.status.name if project.status else '',
        ])
        r = ws.max_row
        ws.cell(row=r, column=2).number_format = 'dd mmm yyyy'
        ws.cell(row=r, column=3).number_format = '0.0%'
        ws.cell(row=r, column=4).number_format = '0.0%'
        ws.cell(row=r, column=5).number_format = 'dd mmm yyyy'
        ws.cell(row=r, column=7).number_format = '#,##0.00'

        # Just the Status cell gets its own colour, same as the small badge
        # it renders as everywhere else in the ERP — a whole-row fill turned
        # into a rainbow once real data (10+ distinct status colours) hit it.
        if project.status and project.status.color:
            hex_color = project.status.color.lstrip('#').upper()
            if len(hex_color) == 6:
                status_cell = ws.cell(row=r, column=11)
                status_cell.fill = PatternFill('solid', fgColor=hex_color)
                status_cell.font = Font(bold=True, color='FFFFFF')

        if (row['delay_months'] or 0) > 0:
            ws.cell(row=r, column=6).fill = delay_fill
        ws.cell(row=r, column=8).font = green_font   # Achieved Milestones
        ws.cell(row=r, column=9).font = red_font     # No. Pending Milestones

    # Weightage of Pending Activities / Completed Activity Weightage: a real
    # Excel colour-scale rule rather than a fixed per-cell fill, so the
    # shading is a genuine gradient across whatever values are in the sheet
    # (and stays correct if someone edits it afterwards).
    if len(rows) > 0:
        first_data_row = header_idx + 1
        last_data_row = ws.max_row
        ws.conditional_formatting.add(
            f'C{first_data_row}:C{last_data_row}',
            ColorScaleRule(start_type='min', start_color='FFFFFF',
                           end_type='max', end_color='C41E3A'))
        ws.conditional_formatting.add(
            f'D{first_data_row}:D{last_data_row}',
            ColorScaleRule(start_type='min', start_color='FFFFFF',
                           end_type='max', end_color='198754'))

    widths = [40, 20, 16, 16, 20, 12, 16, 14, 16, 18, 16]
    for i, w in enumerate(widths[:ncols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _style_export_header_and_grid(ws, header_idx, ncols)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="PM_Dashboard.xlsx"'
    return response


@login_required
@delivery_required
def pm_project_overview(request, pk):
    """Graphical snapshot of one project — status tiles + a milestone chart,
    then cards linking into the dedicated PO / BOQ / matrix pages."""
    project = _pm_visible_project_or_404(request, pk)
    po_document_count = project.documents.filter(document_type='po_document').count()
    costing_sheet_count = project.costing_sheets.count()
    responsibility_matrix = getattr(project, 'responsibility_matrix', None)
    communication_matrix = getattr(project, 'communication_matrix', None)
    has_milestones = project.milestones.exists()
    delivery = board_row(project) if has_milestones else None
    checklist = milestone_checklist(project) if has_milestones else []
    return render(request, 'pmo/pm_project_overview.html', {
        'project': project,
        'po_document_count': po_document_count,
        'costing_sheet_count': costing_sheet_count,
        'responsibility_row_count': len(responsibility_matrix.rows) if responsibility_matrix else 0,
        'communication_row_count': len(communication_matrix.rows) if communication_matrix else 0,
        'delivery': delivery,
        'milestone_checklist': checklist,
        'milestone_rows_json': json.dumps([
            {'label': row['label'], 'variance_display': row['variance_display']}
            for row in project.milestone_display_rows
        ]),
    })


@login_required
@delivery_required
def pm_project_po(request, pk):
    project = _pm_visible_project_or_404(request, pk)
    po_documents = project.documents.filter(document_type='po_document').order_by('-uploaded_at')
    return render(request, 'pmo/pm_project_po.html', {
        'project': project,
        'po_documents': po_documents,
    })


@login_required
@delivery_required
def pm_project_boq(request, pk):
    project = _pm_visible_project_or_404(request, pk)
    costing_sheets = project.costing_sheets.all().order_by('-updated_at')
    return render(request, 'pmo/pm_project_boq.html', {
        'project': project,
        'costing_sheets': costing_sheets,
    })


def _matrix_detail(request, pk, related_name, default_columns, template):
    """Shared by pm_responsibility_matrix and pm_communication_matrix — the
    two grids differ only in which relation/defaults they read."""
    project = _pm_visible_project_or_404(request, pk)
    matrix = getattr(project, related_name, None)
    return render(request, template, {
        'project': project,
        'columns': matrix.columns if matrix else default_columns(),
        'rows': matrix.rows if matrix else [],
    })


@login_required
@delivery_required
def pm_responsibility_matrix(request, pk):
    return _matrix_detail(request, pk, 'responsibility_matrix',
                           default_responsibility_columns, 'pmo/pm_responsibility_matrix.html')


@login_required
@delivery_required
def pm_communication_matrix(request, pk):
    return _matrix_detail(request, pk, 'communication_matrix',
                           default_communication_columns, 'pmo/pm_communication_matrix.html')


def _matrix_edit(request, pk, model, related_name, default_columns, matrix_title,
                  success_message, view_url_name):
    """Shared by pm_responsibility_matrix_edit and pm_communication_matrix_edit
    — same get/parse/sanitize/save shape either way, differing only in which
    model and defaults back it. `matrix` is built in memory rather than
    fetched-or-created: opening this page is a GET and must not itself write
    a blank matrix row."""
    project = _pm_visible_project_or_404(request, pk)
    matrix = getattr(project, related_name, None) or model(project=project)
    if request.method == 'POST':
        try:
            columns_raw = json.loads(request.POST.get('columns_json') or '[]')
            rows_raw = json.loads(request.POST.get('rows_json') or '[]')
        except (ValueError, TypeError):
            columns_raw, rows_raw = [], []
        matrix.columns, matrix.rows = sanitize_grid(columns_raw, rows_raw, default_columns())
        matrix.updated_by = request.user
        matrix.save()
        messages.success(request, success_message)
        return redirect(view_url_name, pk=project.pk)
    return render(request, 'pmo/pm_matrix_edit.html', {
        'project': project,
        'matrix_title': matrix_title,
        'cancel_url_name': view_url_name,
        'columns_json': json.dumps(matrix.columns or default_columns()),
        'rows_json': json.dumps(matrix.rows or []),
    })


@login_required
@delivery_required
def pm_responsibility_matrix_edit(request, pk):
    return _matrix_edit(request, pk, ResponsibilityMatrix, 'responsibility_matrix',
                         default_responsibility_columns, 'Responsibility Matrix',
                         'Responsibility matrix updated.', 'pmo:pm_responsibility_matrix')


@login_required
@delivery_required
def pm_communication_matrix_edit(request, pk):
    return _matrix_edit(request, pk, CommunicationMatrix, 'communication_matrix',
                         default_communication_columns, 'Communication Matrix',
                         'Communication matrix updated.', 'pmo:pm_communication_matrix')


@login_required
@delivery_required
def pm_communication_matrix_export_pdf(request, pk):
    """Client-facing PDF report of the communication matrix — plain table,
    no pricing/internal data, safe to share externally."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    project = _pm_visible_project_or_404(request, pk)
    # Read-only export — never create a row just because someone viewed it.
    matrix = getattr(project, 'communication_matrix', None)
    columns = matrix.columns if matrix else default_communication_columns()
    rows = matrix.rows if matrix else []

    base = getSampleStyleSheet()
    st_title = ParagraphStyle('t', parent=base['Title'], fontName='Helvetica-Bold',
                               fontSize=16, leading=20, textColor=colors.HexColor('#1A1A1A'))
    st_meta = ParagraphStyle('m', parent=base['Normal'], fontName='Helvetica',
                              fontSize=8.5, leading=12, textColor=colors.HexColor('#6C757D'))
    st_head = ParagraphStyle('h', parent=base['Normal'], fontName='Helvetica-Bold',
                              fontSize=8, leading=10, textColor=colors.white)
    st_cell = ParagraphStyle('c', parent=base['Normal'], fontSize=8, leading=10.5)

    buf = BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(
        buf, pagesize=page, topMargin=13 * mm, bottomMargin=13 * mm,
        leftMargin=12 * mm, rightMargin=12 * mm,
        title=f'Communication Matrix - {project.project_name}')
    content_w = page[0] - 24 * mm

    story = [
        Paragraph('Communication Matrix', st_title),
        Spacer(1, 2 * mm),
        Paragraph(f'{project.project_name} &middot; exported {timezone.localtime():%d %B %Y}', st_meta),
        Spacer(1, 5 * mm),
    ]

    col_count = max(len(columns), 1)
    col_width = content_w / col_count
    header = [Paragraph(col.get('name') or '', st_head) for col in columns]
    table_data = [header]
    for row in rows:
        cells = (row or {}).get('cells', {})
        table_data.append([
            Paragraph((cells.get(col['key'], {}) or {}).get('text', '') or '', st_cell)
            for col in columns
        ])
    if len(table_data) == 1:
        table_data.append([Paragraph('&mdash;', st_cell) for _ in columns])

    table = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#C41E3A')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F7F7')]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(table)

    doc.build(story)
    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/pdf')
    filename = f"communication-matrix-{project.serial_number or project.pk}.pdf"
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


# ── Manpower Status ─────────────────────────────────────────────────────────
#
# Every HR employee, office or site, with the extra fields Project
# Management tracks about them filled in where someone has entered them.
# Viewing follows the same audience as the rest of this department
# (can_see_delivery); filling in/clearing details or onboarding someone
# brand new follows the narrower can_update_progress ladder, same as moving
# a milestone figure.

@login_required
@delivery_required
def manpower_list(request):
    from hr.models import Employee

    employees = Employee.objects.select_related('manpower_resource').order_by('full_name')
    q = (request.GET.get('q') or '').strip()
    if q:
        employees = employees.filter(
            Q(full_name__icontains=q)
            | Q(iqama_number__icontains=q)
            | Q(designation__icontains=q)
        )
    return render(request, 'pmo/manpower_list.html', {
        'employees': employees,
        'q': q,
        'can_manage': can_update_progress(request.user),
    })


@login_required
@update_access_required('Only Project Management can add employees.')
def manpower_create(request):
    """Onboard someone who isn't in HR at all yet."""
    if request.method == 'POST':
        form = NewEmployeeManpowerForm(request.POST)
        if form.is_valid():
            resource = form.save(created_by=request.user)
            messages.success(request, f'{resource.employee.full_name} added.')
            return redirect('pmo:manpower_list')
    else:
        form = NewEmployeeManpowerForm()
    return render(request, 'pmo/manpower_new_employee_form.html', {'form': form})


@login_required
@update_access_required('Only Project Management can edit manpower details.')
def manpower_edit(request, employee_pk):
    """Fill in / update the Project-Management fields for one employee who
    already exists in HR."""
    from hr.models import Employee

    employee = get_object_or_404(Employee, pk=employee_pk)
    # Built in memory, not fetched-or-created: opening this page is a GET and
    # must not itself write a row — only an actual save should. Same pattern
    # as the responsibility/communication matrix edit views above.
    resource = getattr(employee, 'manpower_resource', None) or ManpowerResource(employee=employee)
    if request.method == 'POST':
        form = ManpowerDetailsForm(request.POST, instance=resource)
        if form.is_valid():
            form.save()
            messages.success(request, f'{employee.full_name} updated.')
            return redirect('pmo:manpower_list')
    else:
        form = ManpowerDetailsForm(instance=resource)
    return render(request, 'pmo/manpower_details_form.html', {
        'form': form, 'employee': employee,
    })


@require_POST
@login_required
@update_access_required('Only Project Management can clear manpower details.')
def manpower_clear(request, employee_pk):
    """Reset an employee's Project-Management fields back to blank — they
    stay on the list (they're still a real HR employee), just without any
    of the extra details filled in."""
    resource = get_object_or_404(ManpowerResource, employee_id=employee_pk)
    resource.delete()
    messages.success(request, 'Details cleared.')
    return redirect('pmo:manpower_list')


# ── Issue Log ─────────────────────────────────────────────────────────────
#
# The delivery team's risk/issue register, browsed and exported one month
# at a time by date_identified — same "?year=&month=" paging and per-month
# Excel export (plus a zip of every month) as engineer_calendar's calendar.

def _resolve_year_month(request):
    """Read year/month from the query string, falling back to today when
    absent or invalid — same rule engineer_calendar's month paging uses."""
    today = timezone.localdate()
    try:
        year = int(request.GET.get('year', today.year))
        month = int(request.GET.get('month', today.month))
        if not (1 <= month <= 12):
            raise ValueError('Month out of range')
        calendar.monthrange(year, month)
    except (ValueError, TypeError, OverflowError):
        year, month = today.year, today.month
    return year, month


def _adjacent_month(year, month, delta):
    """delta=-1 for the previous month, +1 for the next, wrapping the year."""
    month += delta
    if month < 1:
        month, year = 12, year - 1
    elif month > 12:
        month, year = 1, year + 1
    return year, month


def _build_issue_log_workbook(year, month, visible_projects):
    """Build one month's Issue Log workbook. Pulled out of the export view so
    the all-months zip can reuse it without duplicating the styling.

    `visible_projects` scopes the issues to what this user may see — the
    same region/ownership ladder as everywhere else in pmo — so an export
    can't be used to pull issues for a project someone couldn't otherwise
    reach."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    issues = (ProjectIssue.objects
              .filter(date_identified__year=year, date_identified__month=month,
                      project__in=visible_projects)
              .select_related('project', 'logged_by')
              .order_by('date_identified', 'pk'))

    columns = [
        'ID', 'Status', 'Priority', 'Description', 'Project Name', 'Owner',
        'Estimated Resolution Date', 'Escalation Needed (Y/N)?', 'Impact', 'Actions',
        'Date Identified', 'Logged By', 'Actual Resolution/Completion Date',
        'Final Resolution & Follow-on Actions',
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = calendar.month_abbr[month].upper()
    ncols = len(columns)

    wrap = Alignment(vertical='top', wrap_text=True)

    # Severity legend, same as the top of the source workbook.
    ws.append(['Issue Severity Description'])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws['A1'].font = Font(bold=True, size=12)
    legend = [
        (ProjectIssue.PRIORITY_CRITICAL, 'Issue will stop project progress.'),
        (ProjectIssue.PRIORITY_HIGH, 'Issue will likely impact budget, schedule or scope.'),
        (ProjectIssue.PRIORITY_MEDIUM, 'Issue impacts the project, but could be mitigated to '
                                       'avoid an impact on budget, schedule or scope.'),
        (ProjectIssue.PRIORITY_LOW, 'Issue is low impact and/or low effort to resolve.'),
    ]
    priority_labels = dict(ProjectIssue.PRIORITY_CHOICES)
    for key, desc in legend:
        ws.append([priority_labels[key], desc])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncols)
        ws.cell(row=r, column=1).fill = PatternFill('solid', fgColor=ProjectIssue.PRIORITY_COLORS[key])
        ws.cell(row=r, column=1).font = Font(bold=True, color='FFFFFF')
    ws.append([])

    header_idx = ws.max_row + 1
    ws.append(columns)

    for issue in issues:
        ws.append([
            issue.pk,
            issue.get_status_display(),
            issue.get_priority_display(),
            issue.description,
            issue.project.project_name,
            issue.owner,
            issue.estimated_resolution_date,
            'Y' if issue.escalation_needed else 'N',
            issue.impact,
            issue.actions,
            issue.date_identified,
            issue.logged_by.get_full_name() if issue.logged_by else '',
            issue.actual_resolution_date,
            issue.final_resolution,
        ])
        r = ws.max_row
        for col_idx in (7, 11, 13):
            ws.cell(row=r, column=col_idx).number_format = 'dd mmm yyyy'
        for c in ws[r]:
            c.alignment = wrap
        # Priority colours the whole cell (matching the legend); Status text
        # is green once Closed, the brand red while still Open.
        ws.cell(row=r, column=3).fill = PatternFill('solid', fgColor=issue.priority_color)
        ws.cell(row=r, column=3).font = Font(bold=True, color='FFFFFF')
        status_color = '198754' if issue.status == ProjectIssue.STATUS_CLOSED else 'C41E3A'
        ws.cell(row=r, column=2).font = Font(bold=True, color=status_color)

    widths = [6, 10, 10, 40, 32, 14, 18, 14, 30, 30, 15, 16, 20, 40]
    for i, w in enumerate(widths[:ncols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _style_export_header_and_grid(ws, header_idx, ncols)

    filename = f'Issue_Log_{calendar.month_name[month]}_{year}.xlsx'
    return wb, filename


@login_required
@delivery_required
def issue_log_list(request):
    year, month = _resolve_year_month(request)
    issues = (ProjectIssue.objects
              .filter(date_identified__year=year, date_identified__month=month,
                      project__in=projects_visible_to(request.user))
              .select_related('project', 'logged_by'))
    prev_year, prev_month = _adjacent_month(year, month, -1)
    next_year, next_month = _adjacent_month(year, month, 1)

    return render(request, 'pmo/issue_log_list.html', {
        'issues': issues,
        'year': year,
        'month': month,
        'month_name': calendar.month_name[month],
        'prev_year': prev_year, 'prev_month': prev_month,
        'next_year': next_year, 'next_month': next_month,
        'can_manage': can_update_progress(request.user),
    })


@login_required
@update_access_required('Only Project Management can log issues.')
def issue_log_create(request):
    if request.method == 'POST':
        form = ProjectIssueForm(request.POST, user=request.user)
        if form.is_valid():
            issue = form.save(commit=False)
            issue.logged_by = request.user
            issue.save()
            messages.success(request, 'Issue logged.')
            return redirect(
                f"{reverse('pmo:issue_log_list')}?year={issue.date_identified.year}&month={issue.date_identified.month}")
    else:
        form = ProjectIssueForm(user=request.user, initial={'date_identified': timezone.localdate()})
    return render(request, 'pmo/issue_log_form.html', {'form': form, 'is_new': True})


@login_required
@update_access_required('Only Project Management can edit issues.')
def issue_log_edit(request, pk):
    issue = get_object_or_404(
        ProjectIssue.objects.filter(project__in=projects_visible_to(request.user)), pk=pk)
    if request.method == 'POST':
        form = ProjectIssueForm(request.POST, instance=issue, user=request.user)
        if form.is_valid():
            form.save()
            messages.success(request, 'Issue updated.')
            return redirect(
                f"{reverse('pmo:issue_log_list')}?year={issue.date_identified.year}&month={issue.date_identified.month}")
    else:
        form = ProjectIssueForm(instance=issue, user=request.user)
    return render(request, 'pmo/issue_log_form.html', {
        'form': form, 'is_new': False, 'issue': issue,
    })


@login_required
@delivery_required
def issue_log_export_excel(request):
    year, month = _resolve_year_month(request)
    wb, filename = _build_issue_log_workbook(year, month, projects_visible_to(request.user))
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
@delivery_required
def issue_log_export_all_zip(request):
    """One zip with a full Issue Log workbook for every month that has any
    issues at all — same format as the single-month Export to Excel button."""
    visible_projects = projects_visible_to(request.user)
    months = list(
        ProjectIssue.objects.filter(project__in=visible_projects)
        .values_list('date_identified__year', 'date_identified__month')
        .distinct().order_by('-date_identified__year', '-date_identified__month')
    )
    if not months:
        messages.error(request, 'There are no logged issues to download yet.')
        return redirect('pmo:issue_log_list')

    buf = BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for yr, mo in months:
            month_wb, month_filename = _build_issue_log_workbook(yr, mo, visible_projects)
            month_buf = BytesIO()
            month_wb.save(month_buf)
            zf.writestr(month_filename, month_buf.getvalue())

    response = HttpResponse(buf.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = 'attachment; filename="Issue_Log_All_Months.zip"'
    return response
