"""Excel export of the budget procurement tracker."""
import io
from decimal import Decimal

import openpyxl
from django.test import TestCase
from django.urls import reverse

from accounts.models import User, Role


def _num(value):
    return Decimal(str(value))


class TrackerExcelExportTests(TestCase):
    """The tracker's Export Excel button downloads the same figures the page
    shows, under the same access rules."""

    def setUp(self):
        from projects.models import Region, ProjectStatus, Project
        from costing.models import CostingSheet, CostingSection, CostingLineItem
        self.proc_role, _ = Role.objects.get_or_create(name=Role.PROCUREMENT_MGR)
        self.region = Region.objects.create(name='Arabia')
        self.won = ProjectStatus.objects.create(name='Won', category='won')
        self.project = Project.objects.create(
            project_name='P', status=self.won, region=self.region)
        self.sheet = CostingSheet.objects.create(
            title='S', project=self.project, margin=Decimal('30'),
            discount_rate=Decimal('0'), workflow_stage='finance_approved')
        self.sec = CostingSection.objects.create(
            costing_sheet=self.sheet, section_number='A.1', title='CCTV', order=0)
        self.item = CostingLineItem.objects.create(
            section=self.sec, item_number='1', description='Cam', quantity=Decimal('2'),
            unit='EA', make='Bosch', model_number='X1', vendor_name='Acme',
            base_unit_cost=Decimal('100'), supplier_currency='SAR')
        self.proc = User.objects.create_user(
            'proc', password='pw', role=self.proc_role, region=self.region)
        self.client.force_login(self.proc)
        self.tracker_url = reverse('procurement:bom_procurement_tracker',
                                   kwargs={'sheet_pk': self.sheet.pk})
        self.export_url = reverse('procurement:bom_tracker_export_excel',
                                  kwargs={'sheet_pk': self.sheet.pk})

    def export(self):
        resp = self.client.get(self.export_url)
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.content)).active

    def item_row(self, ws):
        # Row 11 is the header, 12 the A.1 section, 13 the line itself.
        return [c.value for c in ws[13]]

    def order(self, qty):
        from procurement.models import PurchaseOrder
        self.client.post(self.tracker_url, {'item_ids': [str(self.item.pk)],
                                            f'qty_{self.item.pk}': str(qty)})
        return PurchaseOrder.objects.filter(project=self.project).latest('id')

    def test_downloads_an_excel_file(self):
        resp = self.client.get(self.export_url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('attachment; filename="budget-tracker-s-', resp['Content-Disposition'])
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        self.assertEqual(ws['A1'].value, 'S')
        self.assertEqual(ws['A11'].value, 'Item No')
        self.assertEqual(ws['A12'].value, 'A.1 \u00b7 CCTV')

    def test_line_matches_the_tracker(self):
        row = self.item_row(self.export())
        self.assertEqual(row[:5], ['1', 'Cam', 'Bosch \u00b7 X1', 'Acme', 'EA'])
        self.assertEqual([_num(v) for v in row[5:8]], [Decimal('2'), Decimal('0'), Decimal('2')])
        self.item.refresh_from_db()
        self.assertAlmostEqual(float(row[8]), float(self.item.budget_unit_price()), places=2)
        self.assertAlmostEqual(float(row[9]), float(self.item.budget_line_price()), places=2)
        self.assertEqual(row[10], 'Available')
        self.assertIn(row[11], (None, ''))

    def test_partial_order_shows_status_and_po(self):
        po = self.order(1)
        row = self.item_row(self.export())
        self.assertEqual([_num(v) for v in row[6:8]], [Decimal('1'), Decimal('1')])
        self.assertEqual(row[10], 'Partial \u2014 1 of 2 ordered')
        self.assertEqual(row[11], f'{po.po_number} (1)')

    def test_cancelled_po_is_marked_and_not_counted(self):
        from procurement.models import PurchaseOrder
        po = self.order(2)
        PurchaseOrder.objects.filter(pk=po.pk).update(status='cancelled')
        row = self.item_row(self.export())
        self.assertEqual([_num(v) for v in row[6:8]], [Decimal('0'), Decimal('2')])
        self.assertEqual(row[10], 'Available')
        self.assertEqual(row[11], f'{po.po_number} (2) cancelled')

    def test_summary_and_total_match_the_page(self):
        ws = self.export()
        self.item.refresh_from_db()
        total = float(self.item.budget_line_price())
        self.assertEqual(ws['A6'].value, 'Supply Items')
        self.assertEqual([ws['B6'].value, ws['B7'].value, ws['B8'].value], [1, 1, 0])
        self.assertEqual(ws['A9'].value, 'Budget Total (SAR)')
        self.assertAlmostEqual(float(ws['B9'].value), total, places=2)
        self.assertEqual(ws['I14'].value, 'Budget Total (SAR)')
        self.assertAlmostEqual(float(ws['J14'].value), total, places=2)

    def test_another_regions_procurement_is_turned_away(self):
        from projects.models import Region
        other = User.objects.create_user('west', password='pw', role=self.proc_role,
                                         region=Region.objects.create(name='West', code='WEST'))
        self.client.force_login(other)
        resp = self.client.get(self.export_url)
        self.assertRedirects(resp, reverse('procurement:approved_budgets'),
                             fetch_redirect_response=False)

    def test_a_budget_not_yet_approved_is_not_exported(self):
        self.sheet.workflow_stage = 'finance_review'
        self.sheet.save(update_fields=['workflow_stage'])
        resp = self.client.get(self.export_url)
        self.assertRedirects(resp, reverse('procurement:approved_budgets'),
                             fetch_redirect_response=False)

    def test_tracker_page_has_the_button(self):
        self.assertContains(self.client.get(self.tracker_url), self.export_url)
