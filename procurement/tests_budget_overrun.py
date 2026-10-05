"""Over-budget purchase orders: flagging, approval, display and auth.

Over budget is measured in money at two levels, a budget line and the
budget as a whole. A PO is never blocked: it carries on through the normal
approval chain while a BudgetOverrun collects SCM and PM remarks and a COO
decision. One request per PO. A COO rejection sends the PO back to draft
with its approvals cleared; an approval is shown beside the budget, never
added into it.
"""
import io
from datetime import date
from decimal import Decimal as D
from unittest import mock

import openpyxl
from django.contrib.messages import get_messages
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import Role, User
from procurement import budget_overrun as bo
from procurement.models import (
    BudgetOverrun, BudgetOverrunEvent, POStageApprover, POStatusChange, PurchaseOrder)


def a_region(code, name):
    from projects.models import Region
    region, _ = Region.objects.get_or_create(code=code, defaults={'name': name})
    return region


def a_user(username, role_name, **kw):
    role, _ = Role.objects.get_or_create(name=role_name)
    return User.objects.create_user(username, password='pw', role=role, **kw)


def grant(user, codename, allowed=True):
    user.role.permissions.update_or_create(codename=codename, defaults={'allowed': allowed})
    return User.objects.get(pk=user.pk)   # fresh instance: capabilities are cached


class Fixture(TestCase):
    """A won project with one finance-approved budget:
    Camera 10 units, budget 2,000; Cable 100 units, budget 500. Total 2,500."""

    def setUp(self):
        from costing.models import CostingSection, CostingSheet
        from projects.models import Project, ProjectStatus
        self.east = a_region('EAST', 'Eastern')
        self.west = a_region('WEST', 'Western')
        self.won, _ = ProjectStatus.objects.get_or_create(name='Won', defaults={'category': 'won'})
        self.project = Project.objects.create(
            project_name='Jubail', status=self.won, region=self.east, proposal_reference='REF-OVR-1')
        self.sheet = CostingSheet.objects.create(
            title='Jubail CCTV', project=self.project, margin=D('30'),
            discount_rate=D('0'), workflow_stage='finance_approved')
        self.section = CostingSection.objects.create(
            costing_sheet=self.sheet, section_number='A.1', title='Supply', order=0)
        self.camera = self.line('1', 'Camera', '10', '2000')
        self.cable = self.line('2', 'Cable', '100', '500')
        self.boss = a_user('boss', Role.SUPER_ADMIN)
        self._n = 0

    def line(self, no, desc, qty, budget, section=None, **kw):
        from costing.models import CostingLineItem
        return CostingLineItem.all_objects.create(
            section=section or self.section, item_number=no, description=desc,
            quantity=D(qty), unit='EA', base_unit_cost=D('1'), supplier_currency='SAR',
            budget_price=D(budget), **kw)

    def po(self, *lines, status='draft', **kw):
        self._n += 1
        values = dict(po_date=date(2026, 1, 1), po_number=f'PO-OVR-{self._n}', vendor_name='ACME',
                      po_issued_by='Tester', project=self.project,
                      project_name=self.project.project_name, status=status)
        values.update(kw)
        p = PurchaseOrder.objects.create(**values)
        for i, (ln, qty, rate) in enumerate(lines, 1):
            p.items.create(serial_number=i, description=ln.description if ln else 'Unbudgeted',
                           quantity=D(qty), rate_per_unit=D(rate), uom='EA', source_bom_item=ln)
        return p

    def flagged(self, rate='250', **kw):
        """A PO with Camera 10 x rate: at 250 the line is 500 over."""
        p = self.po((self.camera, '10', rate), **kw)
        bo.record_for_po(p)
        return p, p.budget_overruns.get()

    def signers(self):
        self.scm = grant(a_user('scm', Role.PROCUREMENT_MGR, region=self.east), 'po.approve')
        self.pm = grant(a_user('pm', Role.PROJECT_MANAGER, region=self.east), 'po.approve')
        self.coo = grant(a_user('coo', Role.ADMIN, region=self.east), 'po.approve')

    def to_coo(self, o, user=None):
        bo.add_remarks(o, 'scm', user or self.boss, 'Vendor price rose.')
        return bo.add_remarks(o, 'pm', user or self.boss, 'Needed for handover.')


# ─── 1. The calculation ──────────────────────────────────────────────────────

class CalculationTests(Fixture):

    def test_within_budget_nothing_is_over(self):
        self.po((self.camera, '5', '200'))
        r = bo.check_sheet(self.sheet)
        self.assertEqual(r['over'], 0)
        self.assertEqual(r['lines'], [])

    def test_a_line_can_be_over_while_the_budget_is_not(self):
        self.po((self.camera, '10', '250'))
        r = bo.check_sheet(self.sheet)
        self.assertEqual(r['over'], 0)
        self.assertEqual([(l['line'], l['over']) for l in r['lines']], [(self.camera, D('500'))])

    def test_the_whole_budget_over(self):
        self.po((self.camera, '10', '250'), (self.cable, '100', '6'))
        r = bo.check_sheet(self.sheet)
        self.assertEqual((r['budget'], r['spend'], r['over']), (D('2500'), D('3100'), D('600')))
        self.assertEqual({l['line'].pk: l['over'] for l in r['lines']},
                         {self.camera.pk: D('500'), self.cable.pk: D('100')})

    def test_a_cancelled_po_counts_for_nothing(self):
        self.po((self.camera, '10', '900'), status='cancelled')
        self.assertEqual(bo.check_sheet(self.sheet)['lines'], [])

    def test_a_draft_counts(self):
        self.po((self.camera, '10', '300'), status='draft')
        self.assertEqual(len(bo.check_sheet(self.sheet)['lines']), 1)

    def test_the_po_discount_is_applied(self):
        self.po((self.camera, '10', '240'), discount_rate=D('20'))   # 2,400 less 20% = 1,920
        self.assertEqual(bo.check_sheet(self.sheet)['lines'], [])

    def test_vat_is_ignored_because_budgets_carry_none(self):
        self.po((self.camera, '10', '200'), vat_rate=D('15'))         # 2,000 + VAT is still 2,000
        self.assertEqual(bo.check_sheet(self.sheet)['lines'], [])

    def test_a_sub_item_counts_against_its_parent(self):
        bracket = self.line('1.1', 'Bracket', '1', '0', parent_item=self.camera,
                            added_by_procurement=True)
        self.po((self.camera, '10', '200'), (bracket, '1', '100'))
        r = bo.check_sheet(self.sheet)
        self.assertEqual([(l['line'], l['over']) for l in r['lines']], [(self.camera, D('100'))])

    def test_a_sub_item_never_moves_the_approved_total(self):
        self.line('1.1', 'Bracket', '1', '999', parent_item=self.camera, added_by_procurement=True)
        self.assertEqual(bo.check_sheet(self.sheet)['budget'], D('2500'))

    def test_an_optional_section_is_not_budget(self):
        from costing.models import CostingSection
        optional = CostingSection.objects.create(costing_sheet=self.sheet, section_number='A.9',
                                                 title='Optional', order=9, is_optional=True)
        extra = self.line('9', 'Spare', '1', '1000', section=optional)
        self.po((extra, '1', '5000'))
        r = bo.check_sheet(self.sheet)
        self.assertEqual((r['budget'], r['over'], r['lines']), (D('2500'), 0, []))

    def test_an_unlinked_line_on_the_same_po_is_not_counted(self):
        self.po((self.camera, '5', '200'), (None, '1', '99999'))
        self.assertEqual(bo.check_sheet(self.sheet)['spend'], D('1000'))


# ─── 2. Recording requests ───────────────────────────────────────────────────

class RecordingTests(Fixture):

    def test_flagged_once_with_its_line_and_an_event(self):
        p, o = self.flagged()
        self.assertEqual(o.status, BudgetOverrun.AWAITING_REMARKS)
        self.assertEqual([(l.budget_line, l.over) for l in o.lines.all()], [(self.camera, D('500'))])
        self.assertEqual(list(o.events.values_list('action', flat=True)), ['flagged'])
        bo.record_for_po(p)                         # nothing changed: nothing new
        self.assertEqual(p.budget_overruns.count(), 1)
        self.assertEqual(o.events.count(), 1)

    def test_a_change_in_figures_updates_the_same_request(self):
        p, o = self.flagged()
        p.items.update(rate_per_unit=D('300'))
        bo.record_for_po(p)
        o.refresh_from_db()
        self.assertEqual(p.budget_overruns.count(), 1)
        self.assertEqual(o.lines.get().over, D('1000'))
        self.assertEqual(o.events.last().action, 'updated')

    def test_cancelling_resolves_it(self):
        p, o = self.flagged()
        p.record_status_change(to_status='cancelled', changed_by=self.boss)
        o.refresh_from_db()
        self.assertEqual(o.status, BudgetOverrun.RESOLVED)
        self.assertEqual(o.events.last().action, 'resolved')
        self.assertEqual(bo.line_flags(self.sheet), {})

    def test_unlinking_resolves_it(self):
        p, o = self.flagged()
        p.items.update(source_bom_item=None)
        bo.record_for_po(p)
        o.refresh_from_db()
        self.assertEqual(o.status, BudgetOverrun.RESOLVED)

    def test_back_within_budget_resolves_it(self):
        p, o = self.flagged()
        p.items.update(rate_per_unit=D('200'))
        bo.record_for_po(p)
        o.refresh_from_db()
        self.assertEqual(o.status, BudgetOverrun.RESOLVED)

    def test_an_approval_covers_the_same_or_a_smaller_overrun(self):
        p, o = self.flagged()
        bo.decide(self.to_coo(o), self.boss, True)
        bo.record_for_po(p)
        p.items.update(rate_per_unit=D('240'))      # still over, but by less
        bo.record_for_po(p)
        self.assertEqual(p.budget_overruns.count(), 1)

    def test_an_overrun_that_grows_after_approval_needs_a_new_one(self):
        p, o = self.flagged()
        bo.decide(self.to_coo(o), self.boss, True)
        p.items.update(rate_per_unit=D('300'))
        bo.record_for_po(p)
        self.assertEqual(p.budget_overruns.count(), 2)
        self.assertEqual(p.budget_overruns.filter(status=BudgetOverrun.AWAITING_REMARKS).count(), 1)

    def test_after_a_rejection_nothing_reopens_until_the_figures_change(self):
        p, o = self.flagged()
        bo.decide(self.to_coo(o), self.boss, False, 'Too expensive.')
        bo.record_for_po(p)
        self.assertEqual(p.budget_overruns.count(), 1)
        p.items.update(rate_per_unit=D('260'))      # revised, still over
        bo.record_for_po(p)
        self.assertEqual(p.budget_overruns.filter(status=BudgetOverrun.AWAITING_REMARKS).count(), 1)

    def test_one_request_per_po_on_the_same_line(self):
        a = self.po((self.camera, '10', '250'))
        b = self.po((self.camera, '1', '100'))
        bo.record_for_po(a)
        bo.record_for_po(b)
        self.assertEqual(a.budget_overruns.count(), 1)
        self.assertEqual(b.budget_overruns.count(), 1)

    def test_a_po_within_budget_when_saved_is_not_flagged_by_a_later_one(self):
        a = self.po((self.camera, '5', '200'))
        bo.record_for_po(a)
        b = self.po((self.camera, '5', '300'))
        bo.record_for_po(b)
        self.assertFalse(a.budget_overruns.exists())
        self.assertTrue(b.budget_overruns.exists())

    def test_a_failing_check_never_breaks_the_save_that_called_it(self):
        p = self.po((self.camera, '10', '250'))
        with mock.patch('procurement.budget_overrun.record_for_po', side_effect=RuntimeError('boom')):
            with self.assertLogs('procurement.budget_overrun', level='ERROR'):
                self.assertEqual(bo.refresh_safely(p), [])

    def test_the_largest_approved_overrun_is_shown_not_the_sum(self):
        a, oa = self.flagged()                       # Camera 500 over
        bo.decide(self.to_coo(oa), self.boss, True)
        b = self.po((self.camera, '1', '100'))        # Camera now 600 over
        bo.record_for_po(b)
        bo.decide(self.to_coo(b.budget_overruns.get()), self.boss, True)
        self.assertEqual(bo.approved_overruns(self.sheet), {self.camera.pk: D('600')})


# ─── 3. The approval flow ────────────────────────────────────────────────────

class ApprovalFlowTests(Fixture):

    def test_remarks_then_the_coo_decides(self):
        p, o = self.flagged()
        o = bo.add_remarks(o, 'scm', self.boss, 'Price rose.')
        self.assertEqual(o.status, BudgetOverrun.AWAITING_REMARKS)
        o = bo.add_remarks(o, 'pm', self.boss, 'Needed.')
        self.assertEqual(o.status, BudgetOverrun.AWAITING_COO)
        o = bo.decide(o, self.boss, True, 'Fine.')
        self.assertEqual(o.status, BudgetOverrun.APPROVED)
        self.assertEqual(list(o.events.values_list('action', flat=True)),
                         ['flagged', 'scm_remarks', 'pm_remarks', 'approved'])

    def test_the_coo_cannot_decide_before_both_remarks(self):
        p, o = self.flagged()
        bo.add_remarks(o, 'scm', self.boss, 'Price rose.')
        with self.assertRaises(ValueError):
            bo.decide(o, self.boss, True)

    def test_a_reason_is_required_to_reject(self):
        p, o = self.flagged()
        with self.assertRaises(ValueError):
            bo.decide(self.to_coo(o), self.boss, False, '   ')

    def test_empty_remarks_are_refused(self):
        p, o = self.flagged()
        with self.assertRaises(ValueError):
            bo.add_remarks(o, 'scm', self.boss, '  ')

    def test_an_unknown_stage_is_refused(self):
        p, o = self.flagged()
        with self.assertRaises(ValueError):
            bo.add_remarks(o, 'coo', self.boss, 'Not a remarks stage.')

    def test_a_decided_request_cannot_be_changed(self):
        p, o = self.flagged()
        bo.decide(self.to_coo(o), self.boss, True)
        with self.assertRaises(ValueError):
            bo.add_remarks(o, 'scm', self.boss, 'Late.')
        with self.assertRaises(ValueError):
            bo.decide(o, self.boss, False, 'Changed my mind.')

    def test_rejecting_sends_the_po_back_to_draft_with_approvals_cleared(self):
        p, o = self.flagged(status='issued')
        PurchaseOrder.objects.filter(pk=p.pk).update(scm_approved_at=timezone.now(),
                                                     scm_approved_by=self.boss)
        bo.decide(self.to_coo(o), self.boss, False, 'Too expensive.')
        p.refresh_from_db()
        self.assertEqual(p.status, 'draft')
        self.assertIsNone(p.scm_approved_at)
        self.assertIsNone(p.scm_approved_by)
        change = POStatusChange.objects.get(purchase_order=p, to_status='draft')
        self.assertIn('Approvals cleared: SCM Approval by boss', change.reason)
        self.assertIn('Approvals cleared', o.events.get(action='rejected').remarks)

    def test_rejecting_a_draft_still_clears_its_approvals(self):
        p, o = self.flagged(status='draft')
        PurchaseOrder.objects.filter(pk=p.pk).update(scm_approved_at=timezone.now(),
                                                     scm_approved_by=self.boss)
        bo.decide(self.to_coo(o), self.boss, False, 'Too expensive.')
        p.refresh_from_db()
        self.assertEqual(p.status, 'draft')
        self.assertIsNone(p.scm_approved_at)
        self.assertIn('Approvals cleared', o.events.get(action='rejected').remarks)

    def test_approving_leaves_the_po_and_the_budget_alone(self):
        p, o = self.flagged(status='issued')
        bo.decide(self.to_coo(o), self.boss, True)
        p.refresh_from_db()
        self.camera.refresh_from_db()
        self.assertEqual(p.status, 'issued')
        self.assertEqual(self.camera.budget_price, D('2000'))


# ─── 4. Auth boundaries ──────────────────────────────────────────────────────

class RoleTests(Fixture):

    def setUp(self):
        super().setUp()
        self.signers()

    def test_each_stage_only_by_its_role(self):
        p, o = self.flagged()
        with self.assertRaises(PermissionError):
            bo.add_remarks(o, 'scm', self.pm, 'Not mine.')
        with self.assertRaises(PermissionError):
            bo.add_remarks(o, 'pm', self.scm, 'Not mine.')
        bo.add_remarks(o, 'scm', self.scm, 'Price rose.')
        o = bo.add_remarks(o, 'pm', self.pm, 'Needed.')
        for user in (self.scm, self.pm):
            with self.assertRaises(PermissionError):
                bo.decide(o, user, True)
        self.assertEqual(bo.decide(o, self.coo, True).status, BudgetOverrun.APPROVED)

    def test_the_po_approve_capability_is_required(self):
        p, o = self.flagged()
        pm = grant(self.pm, 'po.approve', allowed=False)
        with self.assertRaises(PermissionError):
            bo.add_remarks(o, 'pm', pm, 'No capability.')

    def test_a_super_admin_can_act_on_every_stage(self):
        p, o = self.flagged()
        self.assertEqual(bo.decide(self.to_coo(o, self.boss), self.boss, True).status,
                         BudgetOverrun.APPROVED)


class ViewAuthTests(Fixture):

    def setUp(self):
        super().setUp()
        self.signers()
        self.p, self.o = self.flagged()

    def remarks_url(self, stage='scm', po=None, overrun=None):
        return reverse('procurement:po_overrun_remarks',
                       args=[(po or self.p).pk, (overrun or self.o).pk, stage])

    def test_anonymous_users_are_sent_to_log_in(self):
        resp = self.client.post(self.remarks_url(), {'remarks': 'x'})
        self.assertEqual(resp.status_code, 302)
        self.o.refresh_from_db()
        self.assertEqual(self.o.scm_remarks, '')

    def test_a_get_changes_nothing(self):
        self.client.force_login(self.scm)
        self.client.get(self.remarks_url(), {'remarks': 'x'})
        self.o.refresh_from_db()
        self.assertEqual(self.o.scm_remarks, '')

    def test_the_right_role_can_post(self):
        self.client.force_login(self.scm)
        self.client.post(self.remarks_url('scm'), {'remarks': 'Price rose.'})
        self.o.refresh_from_db()
        self.assertEqual(self.o.scm_remarks, 'Price rose.')

    def test_the_wrong_role_gets_an_error_and_nothing_changes(self):
        self.client.force_login(self.pm)
        resp = self.client.post(self.remarks_url('scm'), {'remarks': 'Not mine.'}, follow=True)
        self.assertTrue(any('cannot' in str(m).lower() for m in get_messages(resp.wsgi_request)))
        self.o.refresh_from_db()
        self.assertEqual(self.o.scm_remarks, '')

    def test_a_pm_from_another_region_gets_not_found(self):
        west_pm = grant(a_user('pm_west', Role.PROJECT_MANAGER, region=self.west), 'po.approve')
        self.client.force_login(west_pm)
        resp = self.client.post(self.remarks_url('pm'), {'remarks': 'Guessing IDs.'})
        self.assertEqual(resp.status_code, 404)
        self.o.refresh_from_db()
        self.assertEqual(self.o.pm_remarks, '')

    def test_a_coo_from_another_region_cannot_decide(self):
        bo.add_remarks(self.o, 'scm', self.boss, 'a')
        bo.add_remarks(self.o, 'pm', self.boss, 'b')
        west_coo = grant(a_user('coo_west', Role.ADMIN, region=self.west), 'po.approve')
        self.client.force_login(west_coo)
        resp = self.client.post(reverse('procurement:po_overrun_decide', args=[self.p.pk, self.o.pk]),
                                {'decision': 'approve'})
        self.assertEqual(resp.status_code, 404)
        self.o.refresh_from_db()
        self.assertEqual(self.o.status, BudgetOverrun.AWAITING_COO)

    def test_a_request_must_belong_to_the_po_in_the_url(self):
        other = self.po((self.cable, '1', '1'))
        self.client.force_login(self.boss)
        resp = self.client.post(self.remarks_url(po=other), {'remarks': 'x'})
        self.assertEqual(resp.status_code, 404)

    def test_an_unknown_decision_changes_nothing(self):
        bo.add_remarks(self.o, 'scm', self.boss, 'a')
        bo.add_remarks(self.o, 'pm', self.boss, 'b')
        self.client.force_login(self.coo)
        self.client.post(reverse('procurement:po_overrun_decide', args=[self.p.pk, self.o.pk]),
                         {'decision': 'maybe'})
        self.o.refresh_from_db()
        self.assertEqual(self.o.status, BudgetOverrun.AWAITING_COO)


# ─── 5. What triggers the check ──────────────────────────────────────────────

class HookTests(Fixture):

    def test_linking_a_po_flags_it(self):
        proc = grant(a_user('proc', Role.PROCUREMENT_MGR, region=self.east), 'po.approve')
        p = self.po((None, '10', '250'), status='issued')
        item = p.items.get()
        self.client.force_login(proc)
        self.client.post(reverse('procurement:po_link_budget', args=[p.pk]),
                         {'sheet': str(self.sheet.pk), f'line_{item.pk}': str(self.camera.pk)})
        self.assertEqual(p.budget_overruns.get().lines.get().budget_line, self.camera)

    def test_ordering_a_sub_item_from_the_tracker_flags_its_parent(self):
        bracket = self.line('1.1', 'Bracket', '1', '100', parent_item=self.camera,
                            added_by_procurement=True)
        self.po((self.camera, '10', '200'), status='issued')   # Camera exactly on budget
        self.client.force_login(self.boss)
        self.client.post(reverse('procurement:bom_procurement_tracker', args=[self.sheet.pk]),
                         {'item_ids': [str(bracket.pk)], f'qty_{bracket.pk}': '1'})
        new_po = PurchaseOrder.objects.filter(items__source_bom_item=bracket).get()
        self.assertEqual(new_po.budget_overruns.get().lines.get().budget_line, self.camera)


# ─── 6. Display ──────────────────────────────────────────────────────────────

class DisplayTests(Fixture):

    def setUp(self):
        super().setUp()
        self.client.force_login(self.boss)

    def test_the_po_page_shows_the_overrun_card(self):
        p, o = self.flagged()
        resp = self.client.get(reverse('procurement:po_detail', args=[p.pk]))
        self.assertContains(resp, 'Over budget')
        self.assertContains(resp, 'Awaiting SCM and PM remarks')
        self.assertNotContains(resp, 'Not linked to any approved budget')

    def test_an_unlinked_po_shows_the_banner(self):
        p = self.po((None, '1', '100'))
        resp = self.client.get(reverse('procurement:po_detail', args=[p.pk]))
        self.assertContains(resp, 'Not linked to any approved budget')

    def test_the_same_line_over_on_another_po_is_listed(self):
        a, _ = self.flagged()
        b = self.po((self.camera, '1', '100'))
        bo.record_for_po(b)
        resp = self.client.get(reverse('procurement:po_detail', args=[b.pk]))
        self.assertContains(resp, 'Also over on')
        self.assertContains(resp, a.po_number)

    def test_the_list_shows_both_badges(self):
        self.flagged()
        self.po((None, '1', '100'))
        resp = self.client.get(reverse('procurement:po_list'))
        self.assertContains(resp, 'Over budget</span>')
        self.assertContains(resp, 'Not linked to budget')

    def test_a_cancelled_po_is_not_marked_as_unlinked(self):
        self.po((None, '1', '100'), status='cancelled')
        resp = self.client.get(reverse('procurement:po_list'))
        self.assertNotContains(resp, 'Not linked to budget')

    def test_the_tracker_flags_the_line(self):
        self.flagged()
        resp = self.client.get(reverse('procurement:bom_procurement_tracker', args=[self.sheet.pk]))
        self.assertContains(resp, 'awaiting approval')

    def test_an_approved_overrun_shows_beside_the_budget_on_the_tracker(self):
        p, o = self.flagged()
        bo.decide(self.to_coo(o), self.boss, True)
        resp = self.client.get(reverse('procurement:bom_procurement_tracker', args=[self.sheet.pk]))
        self.assertContains(resp, '+ 500.00 approved overrun')
        self.assertContains(resp, '= 2,500.00')

    def test_an_approved_overrun_shows_beside_the_budget_for_finance(self):
        p, o = self.flagged()
        bo.decide(self.to_coo(o), self.boss, True)
        resp = self.client.get(reverse('finance:sheet_budget', args=[self.sheet.pk]))
        self.assertContains(resp, 'Approved Overrun')
        self.assertContains(resp, '+ 500.00 approved overrun')


# ─── 7. The link-to-a-budget prompt and imports ──────────────────────────────

def po_workbook(po_number, project_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'PURCHASE ORDER'
    ws['B1'] = date(2026, 1, 1)
    ws['B2'], ws['B3'], ws['B4'] = po_number, 'Projects', 'ACME'
    ws['F1'], ws['F2'], ws['F3'] = 'Procurement', 'p@example.com', project_name
    for col, head in zip('ABCDEFGHIJ', ['S. No.', 'Make', 'Description', '', '', 'Qty',
                                        'UOM', 'Rate (SAR)', 'Total (SAR)', 'Remarks']):
        ws[f'{col}11'] = head
    ws['A12'], ws['C12'], ws['F12'], ws['G12'], ws['H12'] = 1, 'Monitor', 2, 'EA', 350
    ws['C14'] = 'Total'
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(f'{po_number}.xlsx', buf.getvalue())


class PromptTests(Fixture):

    def setUp(self):
        super().setUp()
        self.client.force_login(self.boss)

    def prompt(self, p):
        return self.client.get(reverse('procurement:po_budget_prompt', args=[p.pk]))

    def test_an_unlinked_po_with_approved_budgets_is_asked(self):
        resp = self.prompt(self.po((None, '1', '100')))
        self.assertContains(resp, 'Jubail CCTV')

    def test_a_linked_po_goes_straight_to_its_page(self):
        p = self.po((self.camera, '1', '100'))
        self.assertRedirects(self.prompt(p), reverse('procurement:po_detail', args=[p.pk]),
                             fetch_redirect_response=False)

    def test_a_po_without_a_project_is_not_asked(self):
        p = self.po((None, '1', '100'), project=None)
        self.assertFalse(bo.needs_budget_prompt(p))

    def test_a_project_without_an_approved_budget_is_not_asked(self):
        self.sheet.workflow_stage = 'finance_review'
        self.sheet.save(update_fields=['workflow_stage'])
        self.assertFalse(bo.needs_budget_prompt(self.po((None, '1', '100'))))

    def test_another_regions_pm_cannot_open_the_prompt(self):
        west_pm = grant(a_user('pm_west', Role.PROJECT_MANAGER, region=self.west), 'po.approve')
        self.client.force_login(west_pm)
        self.assertEqual(self.prompt(self.po((None, '1', '100'))).status_code, 404)

    def test_one_imported_excel_po_is_asked(self):
        resp = self.client.post(reverse('procurement:po_import'),
                                {'excel_file': po_workbook('IMP-1', 'jubail ')})
        p = PurchaseOrder.objects.get(po_number='IMP-1')
        self.assertEqual(p.project, self.project)
        self.assertRedirects(resp, reverse('procurement:po_budget_prompt', args=[p.pk]),
                             fetch_redirect_response=False)

    def test_several_imported_excel_pos_are_listed(self):
        resp = self.client.post(reverse('procurement:po_import'),
                                {'excel_file': [po_workbook('IMP-2', 'Jubail'),
                                                po_workbook('IMP-3', 'Jubail')]}, follow=True)
        text = ' '.join(str(m) for m in get_messages(resp.wsgi_request))
        self.assertIn('Not linked to an approved budget: IMP-2, IMP-3', text)


class ProjectMatchingTests(Fixture):

    def match(self, name):
        from procurement.views import _project_matching_name
        return _project_matching_name(name)

    def test_an_exact_name_matches_ignoring_case_and_spaces(self):
        self.assertEqual(self.match('  JUBAIL '), self.project)

    def test_no_match_attaches_nothing(self):
        self.assertIsNone(self.match('Nowhere'))
        self.assertIsNone(self.match(''))

    def test_two_matches_attach_nothing(self):
        from projects.models import Project
        Project.objects.create(project_name='jubail', status=self.won, region=self.west,
                               proposal_reference='REF-OVR-2')
        self.assertIsNone(self.match('Jubail'))

    def test_a_project_in_the_recycle_bin_is_never_matched(self):
        self.project.is_deleted = True
        self.project.save(update_fields=['is_deleted'])
        self.assertIsNone(self.match('Jubail'))


# ─── 8. Notifications ────────────────────────────────────────────────────────

class NotificationTests(Fixture):

    def setUp(self):
        super().setUp()
        self.signers()
        for stage, user in (('scm', self.scm), ('pm', self.pm), ('coo', self.coo)):
            POStageApprover.objects.create(stage=stage, user=user)

    def notified(self, user, verb_start):
        from notifications.models import Notification
        return Notification.objects.filter(recipient=user, verb__startswith=verb_start).exists()

    def test_flagging_tells_the_scm_and_pm_signers_in_app_only(self):
        with self.captureOnCommitCallbacks(execute=True):
            self.flagged()
        self.assertTrue(self.notified(self.scm, 'Over-budget remarks needed'))
        self.assertTrue(self.notified(self.pm, 'Over-budget remarks needed'))
        self.assertFalse(self.notified(self.coo, 'Over-budget remarks needed'))
        self.assertEqual(len(mail.outbox), 0)

    def test_nothing_is_sent_until_the_save_commits(self):
        from notifications.models import Notification
        with self.captureOnCommitCallbacks() as callbacks:
            self.flagged()
            self.assertEqual(Notification.objects.count(), 0)
        self.assertGreaterEqual(len(callbacks), 1)

    def test_both_remarks_tell_the_coo(self):
        p, o = self.flagged()
        with self.captureOnCommitCallbacks(execute=True):
            self.to_coo(o)
        self.assertTrue(self.notified(self.coo, 'Over-budget decision needed'))

    def test_a_rejection_tells_the_creator_scm_and_pm(self):
        creator = a_user('creator', Role.PROCUREMENT_OFF, region=self.east)
        p, o = self.flagged(created_by=creator)
        self.to_coo(o)
        with self.captureOnCommitCallbacks(execute=True):
            bo.decide(o, self.boss, False, 'Too expensive.')
        for user in (creator, self.scm, self.pm):
            self.assertTrue(self.notified(user, 'Over-budget request rejected'))